# -*- coding: utf-8 -*-
import io, re, sys
from collections import OrderedDict
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from data_uc import PROJECT, OBJECTIVES, DOMAINS, DOMAIN_NAME, ACTORS, UCS, SPECS
from data_rules import BRS, WORKFLOWS, FLOWS, ALGO, ERRORS, CONVENTIONS, RBAC, UAT, PLAN, SPRINTS, REVIEW
from data_screens import SCREENS, SIDEBAR
from data_api import APIS, DTO_EXAMPLES, AI_TOOLS
from data_db import TABLES, ENUMS, UC_TABLES, cs_type, pascal
from naming import tr, tr_row, leftovers


# ---------- chuẩn hóa định danh tiếng Anh (ADR-0001). Giữ nguyên nội dung trích từ DA: mô tả UC, luồng đặc tả, rà soát DA.
DOMAINS = [tr_row(d, (2,)) for d in DOMAINS]
ACTORS = [tr_row(a, (2, 4)) for a in ACTORS]
UCS = [tr_row(u, (10, 11)) for u in UCS]
SPECS = {k: tr_row(v, (5,)) for k, v in SPECS.items()}
BRS = [tr_row(b, (3, 4)) for b in BRS]
WORKFLOWS = [tr_row(w) for w in WORKFLOWS]
FLOWS = [tr_row(f) for f in FLOWS]
ALGO = [tr_row(a, (1, 3), api_cols=(2,)) for a in ALGO]
ERRORS = [tr_row(e) for e in ERRORS]
UAT = [tr_row(u) for u in UAT]
SCREENS = [tr_row(s) for s in SCREENS]
APIS = [tr_row(a, (1, 11), api_cols=(3, 8, 9)) for a in APIS]
DTO_EXAMPLES = [tr_row(d, (1, 4), api_cols=(2, 3)) for d in DTO_EXAMPLES]
AI_TOOLS = [tr_row(a, (1, 4, 5), api_cols=(2, 3)) for a in AI_TOOLS]
REVIEW = [tr_row(r, (4,)) for r in REVIEW]
CONVENTIONS = [tr_row(c) for c in CONVENTIONS]

OUT = sys.argv[1]
SRC = "126_DA_NTTT_NguyenXuanHoang_2214_v1.docx"
VER = "v2.1 · 25/10/2026 · định danh tiếng Anh (ADR-0001)"

NAVY, HEAD, LIGHT, ZEBRA = "134D75", "2E6F95", "EAF2F8", "F7FAFC"
FONT = "Calibri"
thin = Side(style="thin", color="C9D6E0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
TOP = Alignment(vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

wb = Workbook()
wb.remove(wb.active)

def put(ws, r, c, v, formula=False, **style):
    cell = ws.cell(row=r, column=c, value=v)
    if isinstance(v, str) and v.startswith("=") and not formula:
        cell.data_type = "s"
    for k, val in style.items():
        setattr(cell, k, val)
    return cell

def sheet(name, title, subtitle, note, headers, widths, rows, center_cols=(), freeze=None, formula_cols=()):
    ws = wb.create_sheet(name)
    n = len(headers)
    for r, (txt, size, bold, color, fill) in enumerate([
        (title, 14, True, "FFFFFF", NAVY), (subtitle, 10, False, "193147", LIGHT), (note, 9, False, "4A5D6E", LIGHT)], start=1):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n)
        put(ws, r, 1, txt, font=Font(name=FONT, size=size, bold=bold, color=color, italic=(r == 3)),
            fill=PatternFill("solid", fgColor=fill), alignment=Alignment(vertical="center", wrap_text=True))
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[3].height = 30
    for c, h in enumerate(headers, start=1):
        put(ws, 4, c, h, font=Font(name=FONT, size=10, bold=True, color="FFFFFF"),
            fill=PatternFill("solid", fgColor=HEAD), alignment=CENTER, border=BORDER)
    ws.row_dimensions[4].height = 32
    for i, row in enumerate(rows):
        r = 5 + i
        for c, v in enumerate(row, start=1):
            put(ws, r, c, v, formula=(c in formula_cols),
                font=Font(name=FONT, size=10), alignment=CENTER if c in center_cols else TOP, border=BORDER,
                fill=PatternFill("solid", fgColor=ZEBRA) if i % 2 else PatternFill())
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = freeze or ("C5" if n > 6 else "A5")
    if rows:
        ws.auto_filter.ref = f"A4:{get_column_letter(n)}{4 + len(rows)}"
    ws.sheet_view.zoomScale = 90
    return ws

def uc_tokens(s):
    out = []
    for t in re.split(r"[,;]", s or ""):
        t = t.strip()
        m = re.match(r"UC-SPA-(\d+)", t)
        if m: out.append(f"UC-SPA-{int(m.group(1)):02d}")
        elif re.fullmatch(r"\d{1,2}", t): out.append(f"UC-SPA-{int(t):02d}")
    return out

def scr_tokens(s):
    return re.findall(r"SCR-[WM]\d{2}", s or "")

# ---------- derived & checks
API_ROWS = []
for i, a in enumerate(APIS, start=1):
    dom, action, method, path, kind, ucs, scrs, roles, req, resp, http, note = a
    API_ROWS.append(dict(id=f"SPA-API-{i:03d}", dom=dom, action=action, method=method, path=path, kind=kind,
                         ucs=uc_tokens(ucs), scrs=scr_tokens(scrs), scr_raw=scrs, roles=roles, req=req, resp=resp, http=http, note=note))
API_BY_ID = {a["id"]: a for a in API_ROWS}
assert len(API_ROWS) == 116, len(API_ROWS)
expect = {"SPA-API-016": ("PATCH", "/api/v1/users/{id}/status"), "SPA-API-036": ("POST", "/api/v1/me/orders"),
          "SPA-API-037": ("POST", "/api/v1/me/orders/{id}/payments"), "SPA-API-039": ("POST", "/api/v1/payments/webhook"),
          "SPA-API-052": ("POST", "/api/v1/me/appointments/book-now"), "SPA-API-053": ("POST", "/api/v1/me/appointments"),
          "SPA-API-064": ("POST", "/api/v1/appointments/{id}/no-show"), "SPA-API-075": ("PUT", "/api/v1/treatments/{id}"),
          "SPA-API-082": ("POST", "/api/v1/staff/{id}/resign"), "SPA-API-087": ("PATCH", "/api/v1/work-schedules/{id}/status"),
          "SPA-API-097": ("POST", "/api/v1/appointments/{id}/complete"), "SPA-API-101": ("POST", "/api/v1/me/reviews"),
          "SPA-API-104": ("PATCH", "/api/v1/reviews/{id}/visibility"), "SPA-API-107": ("GET", "/api/v1/reports/revenue"),
          "SPA-API-110": ("GET", "/api/v1/me/notifications")}
for k, (m, p) in expect.items():
    assert (API_BY_ID[k]["method"], API_BY_ID[k]["path"]) == (m, p), (k, API_BY_ID[k]["method"], API_BY_ID[k]["path"])
for txt in [w[7] for w in WORKFLOWS] + [f[4] for f in FLOWS]:
    for n in re.findall(r"(?:SPA-API-)?(\d{3})", txt):
        assert f"SPA-API-{n}" in API_BY_ID, txt

UC_CODES = [u[0] for u in UCS]
assert len(UC_CODES) == 34
SCR_BY_ID = {s[0]: s for s in SCREENS}
uc_apis = OrderedDict((c, []) for c in UC_CODES)
uc_scrs = OrderedDict((c, []) for c in UC_CODES)
scr_apis = OrderedDict((s[0], []) for s in SCREENS)
for a in API_ROWS:
    for u in a["ucs"]:
        assert u in uc_apis, (a["id"], u)
        uc_apis[u].append(a["id"])
        for s in a["scrs"]:
            if s not in uc_scrs[u]: uc_scrs[u].append(s)
    for s in a["scrs"]:
        assert s in scr_apis, (a["id"], s)
        scr_apis[s].append(a["id"])
for s in SCREENS:
    for u in uc_tokens(s[13]):
        assert u in uc_scrs, (s[0], u)
        if s[0] not in uc_scrs[u]: uc_scrs[u].append(s[0])
for u in UC_CODES:
    assert u in SPECS and u in UC_TABLES, u
    uc_scrs[u].sort()
for code, _, ucs in OBJECTIVES:
    for u in ucs: assert u in UC_CODES, (code, u)
print("UC không có API:", [u for u in UC_CODES if not uc_apis[u]])
print("UC không có màn hình:", [u for u in UC_CODES if not uc_scrs[u]])
print("Màn hình không có API:", [s for s in scr_apis if not scr_apis[s]])

table_names = [t[1] for t in TABLES]
for u, (main, rel) in UC_TABLES.items():
    for t in [main] + [x.strip() for x in rel.split(",") if x.strip()]:
        assert t in table_names, (u, t)
for t, col, _, _, _ in ENUMS:
    for x in t.split(', '): assert x in table_names, x

n_uc, n_api, n_scr = len(UCS), len(API_ROWS), len(SCREENS)
n_tbl_da = sum(1 for t in TABLES if t[5] == "DA")
LAST = 4 + n_uc

# ================= 00 Tổng quan
ws = sheet("00_Tong_quan", "SPA NHƯ NGUYỆT  |  TỔNG HỢP PHÂN TÍCH NGHIỆP VỤ & THIẾT KẾ HỆ THỐNG",
           f"Nguồn chính: {SRC} (báo cáo ĐATN – Chương 2) · tiến độ: đề cương 126_DC · {VER}",
           "Cột 'Nguồn' phân biệt: DA = lấy từ báo cáo; ĐX = đề xuất bổ sung (cần GVHD/chủ spa chốt). Đặc tả UC giữ nguyên nội dung DA; phần bổ sung tách riêng. "
           "UC/bảng CSDL bám DA; màn hình, API là thiết kế triển khai đề xuất — repo SpaSystem hiện mới có foundation.",
           ["Hạng mục", "Giá trị", "Ghi chú"], [34, 80, 70], [], freeze="A5")
info = [
    ("Tên đề tài", PROJECT["ten_de_tai"], ""),
    ("Sinh viên", PROJECT["sinh_vien"], ""),
    ("Người hướng dẫn", PROJECT["gvhd"], "Bộ môn duyệt: " + PROJECT["bo_mon"]),
    ("Thời gian thực hiện", PROJECT["thoi_gian"], "25_Ke_hoach"),
    ("Công nghệ", PROJECT["cong_nghe"], "DA – Phương pháp nghiên cứu"),
    ("Kiến trúc", PROJECT["kien_truc"], "DA – Mở đầu"),
    ("Bảo mật", PROJECT["bao_mat"], "DA – Mở đầu"),
    ("Chống trùng lịch", PROJECT["chong_trung"], "DA – Mở đầu; chi tiết 09_Thuat_toan_lich"),
    ("Phạm vi AI", "Hỗ trợ cung cấp thông tin và tư vấn dịch vụ — không chẩn đoán hay kết luận y khoa", "DA – Mục tiêu"),
]
stats = [
    ("Nghiệp vụ lớn", "=COUNTA('02_Nghiep_vu'!A5:A100)", ""),
    ("Tác nhân", "=COUNTA('03_Actor'!A5:A100)", "4 actor người dùng + 4 hệ thống ngoài/nền"),
    ("Use case", "=COUNTA('04_Checklist_UC'!B5:B200)", "Theo danh sách actor của DA, gộp UC dùng chung"),
    ("UC có đặc tả trong DA", "=COUNTIF('04_Checklist_UC'!H5:H200,\"DA: có đặc tả\")", "Còn lại: đặc tả đề xuất"),
    ("Business rules", "=COUNTA('06_Business_Rules'!A5:A200)", ""),
    ("Màn hình (web + Android)", "=COUNTA('11_Man_hinh'!A5:A200)", "Tách web/Android ở 01_Dashboard"),
    ("API endpoints", "=COUNTA('12_API_Endpoints'!A5:A400)", "SPA-API-001 → %03d" % n_api),
    ("Bảng CSDL theo DA", "=COUNTIF('17_DB_Bang'!F5:F200,\"DA\")", f"+ {len(TABLES)-n_tbl_da} bảng đề xuất bổ sung (ĐX)"),
    ("Vấn đề rà soát trong DA", "=COUNTA('26_Ra_soat_DA'!A5:A200)", "Cần sửa trước khi nộp báo cáo"),
    ("Test case UAT", "=COUNTA('23_UAT'!A5:A200)", ""),
]
r = 5
for k, v, n in info:
    put(ws, r, 1, k, font=Font(name=FONT, bold=True), border=BORDER, alignment=TOP)
    put(ws, r, 2, v, border=BORDER, alignment=TOP); put(ws, r, 3, n, border=BORDER, alignment=TOP); r += 1
r += 1
put(ws, r, 1, "THỐNG KÊ", font=Font(name=FONT, bold=True, color="FFFFFF"), fill=PatternFill("solid", fgColor=HEAD)); r += 1
for k, v, n in stats:
    put(ws, r, 1, k, font=Font(name=FONT, bold=True), border=BORDER, alignment=TOP)
    put(ws, r, 2, v, formula=True, border=BORDER, alignment=Alignment(horizontal="left"))
    put(ws, r, 3, n, border=BORDER, alignment=TOP); r += 1
r += 1
put(ws, r, 1, "DANH MỤC SHEET", font=Font(name=FONT, bold=True, color="FFFFFF"), fill=PatternFill("solid", fgColor=HEAD)); r += 1
SHEET_DESC = [
 ("01_Dashboard","Tiến độ & phân bổ UC theo nghiệp vụ, sprint, trạng thái (công thức tự cập nhật)"),
 ("02_Nghiep_vu","Nghiệp vụ lớn, số UC/API"),
 ("03_Actor","Tác nhân theo DA + hệ thống ngoài"),
 ("04_Checklist_UC","34 UC: ưu tiên, độ khó, sprint, trạng thái BE/FE/Mobile/QA, tiêu chí nghiệm thu"),
 ("05_Dac_ta_UC","Đặc tả UC trích từ DA (Input, Output, Basic/Alternative/Exception flow) + cột rà soát/bổ sung"),
 ("06_Business_Rules","Quy tắc nghiệp vụ, nơi enforce, nguồn"),
 ("07_Workflow","Máy trạng thái: LichHen, DonHang, ThanhToan, quyền sử dụng, gói, liệu trình…"),
 ("08_Quy_trinh","Luồng nghiệp vụ xuyên suốt và đường đi màn hình"),
 ("09_Thuat_toan_lich","Thuật toán giờ trống & chống trùng lịch trên bảng LichLamViec/LichHen"),
 ("10_Sidebar","Menu theo vai trò"),
 ("11_Man_hinh","Màn hình web + Android, route khớp repo"),
 ("12_API_Endpoints","API REST đề xuất"),
 ("13_API_Mau","Mẫu JSON"),
 ("14_AI_Tools","Dữ liệu/tool cho trợ lý AI"),
 ("15_Mapping_UC","UC → màn hình → API → bảng"),
 ("16_Mapping_Man_hinh","Màn hình → API → UC"),
 ("17_DB_Bang","26 bảng theo DA + bảng đề xuất bổ sung"),
 ("18_DB_Cot","Thuộc tính, kiểu, ràng buộc theo DA + kiểu PostgreSQL đề xuất"),
 ("19_DB_FK","Quan hệ khóa ngoại trích từ ràng buộc DA"),
 ("20_Enum","Giá trị đề xuất cho các cột ENUM"),
 ("21_Phan_quyen","Ma trận quyền Khách vãng lai / Khách hàng / Nhân viên / Chủ spa"),
 ("22_Loi_API","Mã lỗi gắn với E1/E2 trong đặc tả"),
 ("23_UAT","Kịch bản kiểm thử"),
 ("24_Quy_uoc_API","Quy ước API"),
 ("25_Ke_hoach","Tiến độ (đề cương) + sprint"),
 ("26_Ra_soat_DA","Lỗi/thiếu sót phát hiện trong file DA và đề xuất sửa"),
 ("27_Truy_vet_MT","Mục tiêu đề tài → UC"),
]
for name, d in SHEET_DESC:
    c = put(ws, r, 1, name, border=BORDER, font=Font(name=FONT, color="0563C1", underline="single"))
    c.hyperlink = f"#'{name}'!A1"
    put(ws, r, 2, d, border=BORDER, alignment=TOP); r += 1
r += 1
put(ws, r, 1, "ĐIỂM CẦN CHỐT", font=Font(name=FONT, bold=True, color="FFFFFF"), fill=PatternFill("solid", fgColor=HEAD)); r += 1
for a, b in [
 ("Trợ lý AI cho khách vãng lai", "DA cho phép nhưng LichSuTroChuyenAI.IdKhachHang not null (RS-11)"),
 ("Xác nhận lịch luồng 2", "Tự xác nhận hay chờ chủ spa xác nhận (RS-12)"),
 ("Thời gian giữ chỗ / hạn thanh toán", "Đề xuất 15 phút"),
 ("Chính sách đổi/hủy lịch", "Đề xuất trước giờ hẹn 24h (RS-16)"),
 ("Điều kiện đánh giá gói", "'Đủ số buổi theo quy định' (RS-15)"),
 ("Hạn dùng dịch vụ lẻ đã mua", "DichVuKhachHang.NgayHetHan null = không giới hạn?"),
 ("Khách không đến", "Có trừ lượt/buổi không"),
 ("Cổng thanh toán", "DA chưa nêu tên (ThanhToan.NhaCungCap)"),
 ("Giao diện", "Tailwind CSS (DA) hay CSS Modules (repo) (RS-02)"),
]:
    put(ws, r, 1, a, font=Font(name=FONT, bold=True), border=BORDER, alignment=TOP)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    put(ws, r, 2, b, border=BORDER, alignment=TOP); r += 1
ws.freeze_panes = None
ws.auto_filter.ref = None

# ================= 02 Nghiệp vụ
rows = []
for i, (code, name, desc, src) in enumerate(DOMAINS, start=1):
    rr = 4 + i
    rows.append((code, name, desc, src,
                 f"=COUNTIF('04_Checklist_UC'!D$5:D$200,A{rr})",
                 f"=COUNTIFS('04_Checklist_UC'!D$5:D$200,A{rr},'04_Checklist_UC'!I$5:I$200,\"P0\")",
                 f"=COUNTIF('12_API_Endpoints'!B$5:B$400,A{rr})",
                 ", ".join(u[0] for u in UCS if u[1] == code)))
sheet("02_Nghiep_vu", "02 · DANH SÁCH NGHIỆP VỤ LỚN", f"{len(DOMAINS)} nghiệp vụ · gom từ danh sách UC của DA",
      "Số UC / P0 / API là công thức đếm từ 04_Checklist_UC và 12_API_Endpoints.",
      ["Mã nhóm", "Nghiệp vụ lớn", "Mô tả phạm vi", "Nguồn", "Số UC", "UC P0", "Số API", "Danh sách UC"],
      [11, 30, 70, 14, 9, 9, 9, 60], rows, center_cols=(1, 5, 6, 7), formula_cols=(5, 6, 7))

# ================= 03 Actor
sheet("03_Actor", "03 · TÁC NHÂN (ACTOR)", "Theo mục 'Liệt kê Actor và Usecase' của DA + hệ thống ngoài", "Mã actor dùng trong cột Actor / Vai trò ở các sheet khác.",
      ["Mã", "Tác nhân", "Phạm vi", "Nền tảng", "Nguồn", "Số UC tham gia"], [10, 24, 90, 22, 40, 14],
      [(a[0], a[1], a[2], a[3], a[4], f"=COUNTIF('04_Checklist_UC'!G$5:G$200,\"*\"&A{5+i}&\"*\")") for i, a in enumerate(ACTORS)],
      center_cols=(1, 6), formula_cols=(6,))

# ================= 04 Checklist UC
STATUS = ["Chưa làm", "Đang làm", "Done chờ Test", "Hoàn thành", "Tạm dừng", "N/A"]
def sc(col, r):
    return f'IF({col}{r}="Hoàn thành",1,IF({col}{r}="Done chờ Test",1,IF({col}{r}="Đang làm",0.5,0)))'
rows = []
for i, u in enumerate(UCS, start=1):
    code, dom, name, desc, actor, src, prio, diff, sprint, plat, acc, note = u
    r = 4 + i
    prog = f'=IFERROR(ROUND(({sc("M",r)}+{sc("N",r)}+{sc("O",r)}+{sc("P",r)})/(4-COUNTIF(M{r}:P{r},"N/A")),4),0)'
    rows.append((i, code, DOMAIN_NAME[dom], dom, name, desc, actor, src, prio, diff, sprint, plat,
                 "Chưa làm", "Chưa làm" if "Web" in plat else "N/A", "Chưa làm" if "Android" in plat else "N/A", "Chưa làm",
                 prog, acc, note))
ws = sheet("04_Checklist_UC", "04 · CHECKLIST USE CASE", f"{n_uc} UC · UC-SPA-01 → UC-SPA-{n_uc:02d} · tên UC theo DA · trạng thái tại 25/10/2026: chưa triển khai (repo mới có foundation)",
           "Chỉnh cột M–P. N/A = UC không có trên nền tảng đó. Tiến độ (Q) = trung bình các cột áp dụng: Đang làm 50%, Done chờ Test/Hoàn thành 100%. "
           "Ưu tiên, độ khó, sprint, tiêu chí nghiệm thu là ĐỀ XUẤT.",
           ["STT", "Mã UC", "Nghiệp vụ lớn", "Mã nhóm", "Tên Use Case (DA)", "Mô tả (DA)", "Actor", "Nguồn", "Ưu tiên", "Độ khó (1-10)",
            "Sprint", "Nền tảng", "Trạng thái BE", "Trạng thái FE Web", "Trạng thái Mobile", "Trạng thái QA", "Tiến độ", "Tiêu chí nghiệm thu (ĐX)", "Ghi chú"],
           [6, 12, 22, 9, 36, 52, 16, 18, 9, 9, 10, 16, 14, 14, 14, 14, 10, 64, 44], rows,
           center_cols=(1, 2, 4, 9, 10, 11, 13, 14, 15, 16, 17), formula_cols=(17,), freeze="F5")
for r in range(5, LAST + 1):
    ws.cell(row=r, column=17).number_format = "0%"
for rng, f1 in [(f"M5:P{LAST}", '"' + ",".join(STATUS) + '"'), (f"I5:I{LAST}", '"P0,P1,P2"'), (f"K5:K{LAST}", '"Sprint 1,Sprint 2,Sprint 3"')]:
    dv = DataValidation(type="list", formula1=f1, allow_blank=True); ws.add_data_validation(dv); dv.add(rng)
for txt, color in [("Hoàn thành", "C6EFCE"), ("Done chờ Test", "DDEBF7"), ("Đang làm", "FFEB9C"), ("Tạm dừng", "F8CBAD"), ("N/A", "EDEDED")]:
    ws.conditional_formatting.add(f"M5:P{LAST}", FormulaRule(formula=[f'M5="{txt}"'], fill=PatternFill("solid", fgColor=color)))
ws.conditional_formatting.add(f"I5:I{LAST}", FormulaRule(formula=['I5="P0"'], font=Font(bold=True, color="C00000")))
ws.conditional_formatting.add(f"H5:H{LAST}", FormulaRule(formula=['LEFT(H5,9)="DA: chỉ l"'], fill=PatternFill("solid", fgColor="FFEB9C")))
ws.conditional_formatting.add(f"Q5:Q{LAST}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="63BE7B"))

# ================= 05 Đặc tả UC
def numbered(lines):
    out, k = [], 0
    for s in lines:
        if re.match(r"^(LUỒNG|\d+\.|ĐX:)", s) or s.isupper():
            out.append(s); continue
        k += 1
        out.append(f"{k}. {s}")
    return "\n".join(out)
rows = []
for i, u in enumerate(UCS, start=1):
    code = u[0]; inp, outp, basic, alt, exc, extra, br = SPECS[code]
    r = 4 + i
    rows.append((i, f"='04_Checklist_UC'!B{r}", f"='04_Checklist_UC'!E{r}", f"='04_Checklist_UC'!F{r}", f"='04_Checklist_UC'!G{r}",
                 "Có trong DA" if u[5].startswith("DA: có") else "ĐX (DA chưa có bảng)",
                 inp, outp, numbered(basic), "\n".join(alt) or "—", "\n".join(exc) or "—", extra or "—", br))
ws = sheet("05_Dac_ta_UC", "05 · ĐẶC TẢ USE CASE", f"{n_uc} đặc tả · cùng thứ tự và mã với 04_Checklist_UC",
           "Cột Input → Exception flow trích nguyên ý từ bảng kịch bản trong DA (đã sửa tên UC bị ghi nhầm). Cột 'Rà soát / bổ sung' là nhận xét và đề xuất, không có trong DA.",
           ["STT", "Mã UC", "Tên Use Case", "Description", "Actors", "Đặc tả", "Input", "Output", "Basic flow", "Alternative flow", "Exception flow", "Rà soát / bổ sung (ĐX)", "Business rules"],
           [6, 12, 30, 36, 16, 14, 30, 30, 70, 50, 50, 50, 16], rows, center_cols=(1, 2, 6), formula_cols=(2, 3, 4, 5), freeze="D5")
ws.conditional_formatting.add(f"F5:F{4+n_uc}", FormulaRule(formula=['LEFT(F5,2)="ĐX"'], fill=PatternFill("solid", fgColor="FFEB9C")))

# ================= 06 BR
sheet("06_Business_Rules", "06 · BUSINESS RULES", f"{len(BRS)} quy tắc", "Cột Nguồn: DA = suy ra trực tiếp từ đặc tả/ràng buộc của DA; ĐX = tham số hoặc quy tắc đề xuất cần chốt.",
      ["Mã BR", "Nhóm", "Tên quy tắc", "Quy tắc chi tiết", "Enforce ở đâu", "UC liên quan", "Nguồn"], [13, 14, 26, 100, 28, 22, 30], BRS, center_cols=(1,))

# ================= 07–09
sheet("07_Workflow", "07 · WORKFLOW TRẠNG THÁI", "Giá trị trạng thái theo 20_Enum (đề xuất vì DA chỉ ghi ENUM)",
      "Mỗi bước kiểm tra trạng thái hiện tại; chuyển sai trạng thái trả 409/400.",
      ["Mã WF", "Đối tượng", "Bước", "Từ trạng thái", "Đến trạng thái", "Actor", "Điều kiện / BR", "API / Trigger"],
      [9, 30, 34, 26, 24, 14, 56, 30], WORKFLOWS, center_cols=(1,))
sheet("08_Quy_trinh", "08 · LUỒNG NGHIỆP VỤ XUYÊN SUỐT", f"{len(FLOWS)} luồng", "Dựa trên mục 'Hoạt động nghiệp vụ' của DA.",
      ["Mã flow", "Tên quy trình", "Các bước nghiệp vụ", "Đường đi màn hình", "API chính", "Rule"], [9, 34, 100, 50, 46, 30], FLOWS, center_cols=(1,))
sheet("09_Thuat_toan_lich", "09 · THUẬT TOÁN GIỜ TRỐNG & CHỐNG TRÙNG LỊCH", "Mục tiêu MT-3 · UC-SPA-15, 16, 17, 23 · bảng LichLamViec, LichHen, NhanVien_DichVu",
      "Khoảng thời gian nửa mở [bắt đầu, kết thúc). NgayHen + Gio hiểu theo giờ Việt Nam.",
      ["Bước", "Tên bước", "Xử lý", "Dữ liệu"], [8, 24, 110, 40], ALGO, center_cols=(1,))

# ================= 10–11
sheet("10_Sidebar", "10 · CẤU TRÚC SIDEBAR", "Menu web theo vai trò", "Icon bộ Lucide; quyền thật kiểm tra ở API.",
      ["Vai trò", "Thứ tự", "Menu", "Icon", "Route", "Mã màn hình"], [12, 8, 28, 20, 30, 14], SIDEBAR, center_cols=(1, 2, 6))
ws = sheet("11_Man_hinh", "11 · DANH MỤC MÀN HÌNH", f"{n_scr} màn hình · Web (React + Vite) và Android (Expo Router)",
           "Trạng thái route: Đã có placeholder = route đã có trong repo (docs/frontend.md, docs/mobile.md); Đề xuất thêm = cần tạo; Có, cần đổi tên = route có nhưng tên chưa khớp DA.",
           ["Mã màn hình", "Nền tảng", "Khu vực / Layout", "Nhóm", "Tên màn hình", "Route", "Trạng thái route", "Mẫu giao diện", "Tabs", "Bộ lọc",
            "Cột bảng / KPI", "Hành động", "Modal / Drawer", "UC liên quan", "Vai trò", "Sprint"],
           [12, 10, 18, 16, 32, 42, 16, 30, 30, 30, 40, 36, 30, 20, 12, 10], SCREENS, center_cols=(1, 2, 7, 16))
ws.conditional_formatting.add(f"G5:G{4+n_scr}", FormulaRule(formula=['G5<>"Đã có placeholder"'], fill=PatternFill("solid", fgColor="FFEB9C")))

# ================= 12 API
rows = [(a["id"], a["dom"], DOMAIN_NAME.get(a["dom"], "Hệ thống"), a["action"], a["method"], a["path"], a["kind"], ", ".join(a["ucs"]) or "—",
         ", ".join(a["scrs"]) or a["scr_raw"], a["roles"], a["req"], a["resp"], a["http"], a["note"],
         "Đã có" if "ĐÃ CÓ" in a["note"] else "Đề xuất") for a in API_ROWS]
ws = sheet("12_API_Endpoints", "12 · DANH MỤC API ENDPOINTS", f"{n_api} endpoint · base /api/v1 · envelope {{success, message, data, errors, traceId}}",
           "Đặc tả đề xuất, chưa triển khai. Tên trường JSON theo thuộc tính DA (camelCase). ALL = mọi người kể cả khách vãng lai.",
           ["ID API", "Mã nhóm", "Nghiệp vụ", "Hành động", "Method", "Endpoint", "Loại", "UC", "Màn hình", "Vai trò", "Request", "Response", "HTTP", "Ghi chú", "Trạng thái"],
           [13, 9, 22, 34, 9, 46, 11, 20, 30, 20, 56, 56, 8, 40, 11], rows, center_cols=(1, 2, 5, 7, 13, 15), freeze="G5")
for m, color in [("GET", "DDEBF7"), ("POST", "C6EFCE"), ("PUT", "FFEB9C"), ("PATCH", "FCE4D6"), ("DELETE", "F8CBAD")]:
    ws.conditional_formatting.add(f"E5:E{4+n_api}", FormulaRule(formula=[f'E5="{m}"'], fill=PatternFill("solid", fgColor=color)))

sheet("13_API_Mau", "13 · MẪU REQUEST / RESPONSE", "JSON minh họa", "ID minh họa; DATETIME dạng UTC.",
      ["API ID", "Mục đích", "Request", "Response", "Lưu ý"], [13, 30, 70, 90, 50], DTO_EXAMPLES, center_cols=(1,))
sheet("14_AI_Tools", "14 · DỮ LIỆU / TOOL CHO TRỢ LÝ AI", "UC-SPA-32 bước 4 'chuẩn bị thông tin dịch vụ liên quan' · BR-AI-01..04",
      "Đề xuất: backend tự truy vấn dữ liệu (hoặc định nghĩa tool cho LLM); IdKhachHang luôn lấy từ JWT.",
      ["Tool", "Mục đích", "Tham số", "Kết quả", "Ai dùng", "Bảng"], [22, 40, 40, 50, 12, 46], AI_TOOLS)

# ================= 15–16 Mapping
rows = []
UC_BY = {u[0]: u for u in UCS}
for u in UC_CODES:
    x = UC_BY[u]; scrs = uc_scrs[u]
    main, rel = UC_TABLES[u]
    rows.append((u, x[2], DOMAIN_NAME[x[1]], x[8], x[6], ", ".join(scrs) or "—", "; ".join(SCR_BY_ID[s][5] for s in scrs) or "—",
                 len(uc_apis[u]), ", ".join(uc_apis[u]), main, rel, "Đủ" if uc_apis[u] and scrs else "Thiếu"))
ws = sheet("15_Mapping_UC", "15 · TRUY VẾT USE CASE → MÀN HÌNH → API → BẢNG", "Sinh tự động từ 11_Man_hinh, 12_API_Endpoints, 17_DB_Bang",
           "Độ phủ = Đủ khi UC có ít nhất 1 màn hình và 1 API.",
           ["Mã UC", "Tên UC", "Nghiệp vụ", "Sprint", "Ưu tiên", "Màn hình", "Route FE", "Số API", "API IDs", "Bảng chính", "Bảng liên quan", "Độ phủ"],
           [12, 36, 22, 10, 8, 30, 50, 8, 60, 22, 50, 9], rows, center_cols=(1, 4, 5, 8, 12))
ws.conditional_formatting.add(f"L5:L{4+len(rows)}", FormulaRule(formula=['L5="Thiếu"'], fill=PatternFill("solid", fgColor="F8CBAD")))
rows = [(s[0], s[1], s[4], s[5], s[6], len(scr_apis[s[0]]), ", ".join(scr_apis[s[0]]) or "—", ", ".join(uc_tokens(s[13])), s[15]) for s in SCREENS]
sheet("16_Mapping_Man_hinh", "16 · TRUY VẾT MÀN HÌNH → API → UC", f"{n_scr} màn hình", "API sử dụng sinh từ cột Màn hình của 12_API_Endpoints.",
      ["Mã màn hình", "Nền tảng", "Tên màn hình", "Route", "Trạng thái route", "Số API", "API sử dụng", "UC", "Sprint"],
      [12, 10, 34, 44, 16, 8, 90, 34, 10], rows, center_cols=(1, 2, 5, 6, 9))

# ================= 17–20 DB (tiếng Anh, ADR-0001)
tbl_rows, col_rows, fk_rows = [], [], []
uc_by_table = OrderedDict((t[1], []) for t in TABLES)
for u, (main, rel) in UC_TABLES.items():
    for t in [main] + [x.strip() for x in rel.split(",") if x.strip()]:
        if u not in uc_by_table[t]: uc_by_table[t].append(u)
for grp, name, entity, da_name, purpose, src, sprint, cols in TABLES:
    nfk = 0
    for c in cols:
        col, pg, null, cons, desc, da_col = c.split("|")
        col_src = "DA" if da_col else "ĐX"
        col_rows.append((grp, name, col, pascal(col), pg, cs_type(pg, cons), "Có" if null == "Y" else "Không", cons, desc, da_col or "—", col_src))
        for m in re.finditer(r"FK → (\w+)\.(\w+)( \(CASCADE\))?", cons):
            nfk += 1
            fk_rows.append((name, col, m.group(1), m.group(2), "Có" if null == "Y" else "Không", "CASCADE" if m.group(3) else "RESTRICT",
                            f"fk_{name}_{m.group(1)}_{col}", col_src))
    tbl_rows.append((grp, name, entity, da_name, purpose, src, len(cols), nfk, sprint, ", ".join(uc_by_table[name]) or "—"))
ws = sheet("17_DB_Bang", "17 · DANH MỤC BẢNG CSDL (POSTGRESQL)", f"{n_tbl_da} bảng theo DA + {len(TABLES)-n_tbl_da} bảng đề xuất bổ sung · tên tiếng Anh snake_case",
           "Tên chuẩn hóa tiếng Anh theo ADR-0001 và docs/domain/glossary.md; cột 'Tên trong DA' để đối chiếu với báo cáo. Bảng nguồn 'ĐX' không có trong DA (RS-17).",
           ["Nhóm", "Bảng", "Entity (C#)", "Tên trong DA", "Mục đích", "Nguồn", "Số cột", "Số FK", "Sprint", "UC liên quan"],
           [9, 26, 22, 22, 56, 9, 8, 8, 10, 64], tbl_rows, center_cols=(1, 6, 7, 8, 9), freeze="C5")
ws.conditional_formatting.add(f"F5:F{4+len(tbl_rows)}", FormulaRule(formula=['LEFT(F5,2)="ĐX"'], fill=PatternFill("solid", fgColor="FFEB9C")))
ws = sheet("18_DB_Cot", "18 · CHI TIẾT CỘT", f"{len(col_rows)} cột · snake_case trong CSDL, PascalCase trong C#",
           "Cột có 'Tên trong DA' = lấy từ báo cáo; '—' = bổ sung (BaseAuditableEntity hoặc bảng đề xuất). Ngày/giờ lịch hẹn và ca làm là giờ Việt Nam (ADR-0003); mọi *_at_utc là UTC.",
           ["Nhóm", "Bảng", "Cột", "Thuộc tính C#", "Kiểu PostgreSQL", "Kiểu C#", "Cho phép null", "Ràng buộc", "Mô tả", "Tên trong DA", "Nguồn"],
           [9, 24, 28, 26, 16, 22, 10, 50, 44, 22, 8], col_rows, center_cols=(1, 7, 11), freeze="D5")
ws.conditional_formatting.add(f"K5:K{4+len(col_rows)}", FormulaRule(formula=['K5="ĐX"'], fill=PatternFill("solid", fgColor="FFEB9C")))
sheet("19_DB_FK", "19 · QUAN HỆ KHÓA NGOẠI", f"{len(fk_rows)} quan hệ", "Mặc định RESTRICT để giữ lịch sử; CASCADE chỉ cho cart_items → carts. Tên ràng buộc theo docs/conventions/naming.md.",
      ["Bảng con", "Cột FK", "Bảng cha", "Khóa cha", "Cho phép null", "Hành vi xóa", "Tên ràng buộc", "Nguồn"], [24, 28, 22, 10, 12, 12, 56, 8], fk_rows, center_cols=(5, 6, 8))
enum_rows = []
for tbl, col, enum, values, note in ENUMS:
    for i, (da, member, db) in enumerate(values):
        enum_rows.append((tbl, col, enum, member, db, da, note if i == 0 else ""))
sheet("20_Enum", "20 · GIÁ TRỊ ENUM", f"{len(ENUMS)} enum · {len(enum_rows)} giá trị", "C# dùng thành viên PascalCase; CSDL và JSON lưu chuỗi snake_case (CHECK). Cột 'Giá trị trong DA' là tên đề xuất ban đầu theo báo cáo (DA chỉ ghi ENUM — RS-18).",
      ["Bảng", "Cột", "Enum C#", "Thành viên C#", "Giá trị lưu (DB/JSON)", "Giá trị trong DA", "Ghi chú"], [26, 14, 24, 22, 22, 18, 60], enum_rows)

# ================= 21–24
ws = sheet("21_Phan_quyen", "21 · MA TRẬN PHÂN QUYỀN", "Theo danh sách actor của DA", "✓ = được phép; ghi chú = giới hạn phạm vi. Kiểm tra tại API (BR-SEC-01).",
           ["Permission", "Mô tả", "UC", "Khách vãng lai", "Khách hàng", "Nhân viên", "Chủ spa"], [28, 56, 22, 22, 12, 24, 16], RBAC, center_cols=(4, 5, 6, 7))
ws.conditional_formatting.add(f"D5:G{4+len(RBAC)}", FormulaRule(formula=['LEFT(D5,1)="✓"'], fill=PatternFill("solid", fgColor="C6EFCE")))
sheet("22_Loi_API", "22 · MÃ LỖI API", f"{len(ERRORS)} mã", "Gắn với các luồng ngoại lệ E1/E2 trong đặc tả; khớp exception middleware của foundation.",
      ["HTTP", "Mã lỗi", "Điều kiện", "Hướng xử lý FE", "Tham chiếu"], [8, 42, 56, 44, 30], ERRORS, center_cols=(1,))
ws = sheet("23_UAT", "23 · KỊCH BẢN KIỂM THỬ", f"{len(UAT)} test case", "Giai đoạn kiểm thử theo đề cương: 23/11 → 06/12/2026. DA nhấn mạnh kiểm thử cơ chế chống trùng lịch.",
           ["Mã", "Nhóm", "UC", "Tình huống", "Kết quả mong đợi", "Ưu tiên", "Kết quả", "Ghi chú"],
           [9, 14, 12, 70, 70, 9, 14, 30], [t + ("Chưa chạy", "") for t in UAT], center_cols=(1, 3, 6, 7))
dv = DataValidation(type="list", formula1='"Chưa chạy,Đạt,Không đạt,Bị chặn"'); ws.add_data_validation(dv); dv.add(f"G5:G{4+len(UAT)}")
for txt, color in [("Đạt", "C6EFCE"), ("Không đạt", "F8CBAD"), ("Bị chặn", "FFEB9C")]:
    ws.conditional_formatting.add(f"G5:G{4+len(UAT)}", FormulaRule(formula=[f'G5="{txt}"'], fill=PatternFill("solid", fgColor=color)))
sheet("24_Quy_uoc_API", "24 · QUY ƯỚC API", "Áp dụng cho web và Android", "",
      ["Hạng mục", "Quy ước", "Ghi chú"], [24, 80, 70], CONVENTIONS)

# ================= 25 Kế hoạch
ws = sheet("25_Ke_hoach", "25 · KẾ HOẠCH THỰC HIỆN", "Bảng tiến độ của đề cương (DA không có) + phân bổ sprint đề xuất", "Tình trạng đối chiếu với repo tại 25/10/2026.",
           ["TT", "Từ ngày", "Đến ngày", "Nội dung công việc", "Kết quả dự kiến", "Phạm vi / tài liệu", "Tình trạng (25/10/2026)"],
           [6, 12, 12, 50, 40, 50, 36], PLAN, center_cols=(1, 2, 3))
r = 5 + len(PLAN) + 1
for c, h in enumerate(["Sprint", "Chủ đề", "Phạm vi", "Thời gian dự kiến", "Số UC", "UC P0"], start=1):
    put(ws, r, c, h, font=Font(name=FONT, bold=True, color="FFFFFF"), fill=PatternFill("solid", fgColor=HEAD), border=BORDER, alignment=CENTER)
for s in SPRINTS:
    r += 1
    for c, v in enumerate(s, start=1):
        put(ws, r, c, v, border=BORDER, alignment=TOP)
    put(ws, r, 5, f"=COUNTIF('04_Checklist_UC'!K$5:K$200,A{r})", formula=True, border=BORDER, alignment=CENTER)
    put(ws, r, 6, f"=COUNTIFS('04_Checklist_UC'!K$5:K$200,A{r},'04_Checklist_UC'!I$5:I$200,\"P0\")", formula=True, border=BORDER, alignment=CENTER)

# ================= 26 Rà soát DA
ws = sheet("26_Ra_soat_DA", "26 · RÀ SOÁT TÀI LIỆU DA", f"{len(REVIEW)} vấn đề phát hiện khi đọc {SRC}",
           "Nên sửa trước khi nộp báo cáo. Cột 'Đã xử lý' để theo dõi.",
           ["Mã", "Vị trí trong DA", "Loại", "Vấn đề", "Đề xuất sửa", "Mức độ", "Đã xử lý"], [8, 34, 16, 70, 70, 10, 12],
           [x + ("Chưa",) for x in REVIEW], center_cols=(1, 6, 7))
dv = DataValidation(type="list", formula1='"Chưa,Đã sửa,Bỏ qua"'); ws.add_data_validation(dv); dv.add(f"G5:G{4+len(REVIEW)}")
ws.conditional_formatting.add(f"F5:F{4+len(REVIEW)}", FormulaRule(formula=['F5="Cao"'], fill=PatternFill("solid", fgColor="F8CBAD")))
ws.conditional_formatting.add(f"G5:G{4+len(REVIEW)}", FormulaRule(formula=['G5="Đã sửa"'], fill=PatternFill("solid", fgColor="C6EFCE")))

# ================= 27 Truy vết mục tiêu
rows = []
for i, (code, text, ucs) in enumerate(OBJECTIVES):
    rr = 5 + i
    rows.append((code, text, ", ".join(ucs), len(ucs), f'=IF(D{rr}>0,"Đã phủ","CHƯA PHỦ")'))
ws = sheet("27_Truy_vet_MT", "27 · TRUY VẾT MỤC TIÊU ĐỀ TÀI → USE CASE", "Mục 'Mục tiêu đề tài' (MT-1..5) và 'Khảo sát người dùng' (KS) của DA",
           "Mỗi mục tiêu phải có ít nhất 1 UC.", ["Mã", "Mục tiêu (DA)", "UC đáp ứng", "Số UC", "Kiểm tra"], [8, 90, 70, 8, 12], rows,
           center_cols=(1, 4, 5), formula_cols=(5,))

# ================= 01 Dashboard
ws = wb.create_sheet("01_Dashboard", 1)
CL = "'04_Checklist_UC'"
ws.merge_cells("B2:I2"); put(ws, 2, 2, "DASHBOARD TIẾN ĐỘ & PHÂN RÃ NGHIỆP VỤ — SPA NHƯ NGUYỆT", font=Font(name=FONT, size=16, bold=True, color="FFFFFF"), fill=PatternFill("solid", fgColor=NAVY))
ws.merge_cells("B3:I3"); put(ws, 3, 2, f"Tự cập nhật từ các sheet · nguồn {SRC} · {VER}", font=Font(name=FONT, italic=True, color="4A5D6E"))
kpis = [("TỔNG USE CASE", f"=COUNTA({CL}!B5:B200)", "0"), ("UC P0", f"=COUNTIF({CL}!I5:I200,\"P0\")", "0"),
        ("API", "=COUNTA('12_API_Endpoints'!A5:A400)", "0"), ("MÀN HÌNH", "=COUNTA('11_Man_hinh'!A5:A200)", "0"),
        ("BẢNG DA", "=COUNTIF('17_DB_Bang'!F5:F200,\"DA\")", "0"), ("VẤN ĐỀ DA", "=COUNTIF('26_Ra_soat_DA'!G5:G200,\"Chưa\")", "0"),
        ("TB TIẾN ĐỘ", f"=IFERROR(AVERAGE({CL}!Q5:Q200),0)", "0%")]
for i, (k, f, fmt) in enumerate(kpis):
    c = 2 + i
    put(ws, 5, c, k, font=Font(name=FONT, size=9, bold=True, color="193147"), fill=PatternFill("solid", fgColor="F0F4F6"), alignment=CENTER, border=BORDER)
    cell = put(ws, 6, c, f, formula=True, font=Font(name=FONT, size=18, bold=True, color=NAVY), alignment=Alignment(horizontal="center"), border=BORDER)
    cell.number_format = fmt
ws.row_dimensions[6].height = 30

def head(r, cols):
    for i, h in enumerate(cols):
        put(ws, r, 2 + i, h, font=Font(name=FONT, bold=True, color="FFFFFF"), fill=PatternFill("solid", fgColor=HEAD), alignment=CENTER, border=BORDER)

r = 8
put(ws, r, 2, "1. PHÂN BỔ THEO NGHIỆP VỤ LỚN", font=Font(name=FONT, bold=True, size=12, color=NAVY)); r += 1
head(r, ["Mã", "Nghiệp vụ lớn", "Số UC", "Tỷ lệ", "UC P0", "Số API", "TB tiến độ", "Trạng thái"])
dom_first = r + 1
for code, name, _, _ in DOMAINS:
    r += 1
    vals = [code, name, f"=COUNTIF({CL}!D$5:D$200,B{r})", f"=IFERROR(D{r}/$D${dom_first+len(DOMAINS)},0)",
            f"=COUNTIFS({CL}!D$5:D$200,B{r},{CL}!I$5:I$200,\"P0\")", f"=COUNTIF('12_API_Endpoints'!B$5:B$400,B{r})",
            f"=IFERROR(AVERAGEIF({CL}!D$5:D$200,B{r},{CL}!Q$5:Q$200),0)",
            f'=IF(D{r}=0,"Không có UC",IF(H{r}=1,"Hoàn thành",IF(H{r}>0,"Đang triển khai","Chưa bắt đầu")))']
    for i, v in enumerate(vals):
        cell = put(ws, r, 2 + i, v, formula=True, border=BORDER, alignment=CENTER if i != 1 else TOP)
        if i in (3, 6): cell.number_format = "0%"
dom_last = r
r += 1
put(ws, r, 2, "TỔNG", font=Font(name=FONT, bold=True), border=BORDER); put(ws, r, 3, "(chưa gồm 3 API hệ thống /system, /health)", border=BORDER, font=Font(name=FONT, italic=True, size=9))
for col, f in [(4, f"=SUM(D{dom_first}:D{dom_last})"), (5, f"=IFERROR(D{r}/D{r},0)"), (6, f"=SUM(F{dom_first}:F{dom_last})"),
               (7, f"=SUM(G{dom_first}:G{dom_last})"), (8, f"=IFERROR(AVERAGE({CL}!Q5:Q200),0)")]:
    cell = put(ws, r, col, f, formula=True, font=Font(name=FONT, bold=True), border=BORDER, alignment=CENTER)
    if col in (5, 8): cell.number_format = "0%"
ws.conditional_formatting.add(f"H{dom_first}:H{dom_last}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="63BE7B"))
chart = BarChart(); chart.type = "bar"; chart.title = "Số UC theo nghiệp vụ"; chart.style = 10
chart.add_data(Reference(ws, min_col=4, min_row=dom_first - 1, max_row=dom_last), titles_from_data=True)
chart.set_categories(Reference(ws, min_col=3, min_row=dom_first, max_row=dom_last))
chart.legend = None; chart.height = 9; chart.width = 16
ws.add_chart(chart, f"K{dom_first-1}")

r += 2
put(ws, r, 2, "2. PHÂN BỔ THEO SPRINT", font=Font(name=FONT, bold=True, size=12, color=NAVY)); r += 1
head(r, ["Sprint", "Chủ đề", "Số UC", "UC P0", "Tổng độ khó", "TB tiến độ", "Thời gian"])
for s in SPRINTS:
    r += 1
    vals = [s[0], s[1], f"=COUNTIF({CL}!K$5:K$200,B{r})", f"=COUNTIFS({CL}!K$5:K$200,B{r},{CL}!I$5:I$200,\"P0\")",
            f"=SUMIF({CL}!K$5:K$200,B{r},{CL}!J$5:J$200)", f"=IFERROR(AVERAGEIF({CL}!K$5:K$200,B{r},{CL}!Q$5:Q$200),0)", s[3]]
    for i, v in enumerate(vals):
        cell = put(ws, r, 2 + i, v, formula=True, border=BORDER, alignment=CENTER if i not in (1, 6) else TOP)
        if i == 5: cell.number_format = "0%"
r += 2
put(ws, r, 2, "3. TRẠNG THÁI TRIỂN KHAI", font=Font(name=FONT, bold=True, size=12, color=NAVY)); r += 1
head(r, ["Trạng thái", "BE", "FE Web", "Mobile", "QA"])
for stt in STATUS:
    r += 1
    put(ws, r, 2, stt, border=BORDER)
    for i, col in enumerate("MNOP"):
        put(ws, r, 3 + i, f'=COUNTIF({CL}!{col}$5:{col}$200,B{r})', formula=True, border=BORDER, alignment=CENTER)
r += 2
put(ws, r, 2, "4. MÀN HÌNH, API, TÀI LIỆU", font=Font(name=FONT, bold=True, size=12, color=NAVY)); r += 1
head(r, ["Chỉ tiêu", "Giá trị"])
for k, f in [("Màn hình Web", "=COUNTIF('11_Man_hinh'!B5:B200,\"Web\")"), ("Màn hình Android", "=COUNTIF('11_Man_hinh'!B5:B200,\"Android\")"),
             ("Route đã có placeholder", "=COUNTIF('11_Man_hinh'!G5:G200,\"Đã có placeholder\")"), ("Route cần tạo / đổi tên", "=COUNTA('11_Man_hinh'!G5:G200)-COUNTIF('11_Man_hinh'!G5:G200,\"Đã có placeholder\")"),
             ("API GET", "=COUNTIF('12_API_Endpoints'!E5:E400,\"GET\")"), ("API ghi (POST/PUT/PATCH/DELETE)", "=COUNTA('12_API_Endpoints'!E5:E400)-COUNTIF('12_API_Endpoints'!E5:E400,\"GET\")"),
             ("API đã có trong repo", "=COUNTIF('12_API_Endpoints'!O5:O400,\"Đã có\")"),
             ("UC chưa có đặc tả trong DA", f"=COUNTIF({CL}!H5:H200,\"DA: chỉ*\")"),
             ("Vấn đề DA mức Cao chưa sửa", "=COUNTIFS('26_Ra_soat_DA'!F5:F200,\"Cao\",'26_Ra_soat_DA'!G5:G200,\"Chưa\")"),
             ("Mục tiêu chưa phủ", "=COUNTIF('27_Truy_vet_MT'!E5:E50,\"CHƯA PHỦ\")")]:
    r += 1
    put(ws, r, 2, k, border=BORDER); put(ws, r, 3, f, formula=True, border=BORDER, alignment=CENTER)
for col, w in zip("ABCDEFGHI", [2, 16, 32, 14, 12, 12, 12, 14, 18]):
    ws.column_dimensions[col].width = w
ws.sheet_view.showGridLines = False
ws.freeze_panes = "A7"

KEEP_DA = {"05_Dac_ta_UC": "DGHIJK", "26_Ra_soat_DA": "BCD", "00_Tong_quan": "ABC", "04_Checklist_UC": "EF", "17_DB_Bang": "D", "18_DB_Cot": "J", "20_Enum": "F", "03_Actor": "E", "27_Truy_vet_MT": "B"}
left = []
for wsx in wb.worksheets:
    keep = KEEP_DA.get(wsx.title, "")
    for row in wsx.iter_rows(min_row=5):
        for cell in row:
            if not hasattr(cell, "column_letter") or cell.column_letter in keep or cell.data_type == "f": continue
            found = leftovers(cell.value)
            if found: left.append((wsx.title, cell.coordinate, sorted(set(found))))
print("Định danh tiếng Việt còn sót:", len(left))
for x in left[:60]: print("  ", x)

wb.calculation.fullCalcOnLoad = True
wb.active = 0
wb.save(OUT)
print("OK:", OUT, "| UC", n_uc, "| API", n_api, "| SCR", n_scr, "| TBL", len(TABLES), f"({n_tbl_da} DA)", "| cột", len(col_rows), "| FK", len(fk_rows), "| RS", len(REVIEW))
